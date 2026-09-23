# -*- coding: utf-8 -*-
"""
Rewrite Healthcare_MEDTHRU_Proposal.docx into plain, MUET Band 5 English.

Edits the existing document in place rather than regenerating it, so the cover
page image, the five figures, table widths, styles and the page-number footer
all survive untouched. Only the words change.

Rules applied to every rewritten sentence:
  * one idea per sentence, under roughly 25 words
  * no em dashes anywhere
  * plain vocabulary in place of academic vocabulary
  * technical accuracy preserved

Run:  python docs/simplify_proposal.py
"""

import copy
import re
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC = r'c:\Users\amiru\Desktop\Healthcare_MEDTHRU_Proposal_original.docx'
OUT = r'c:\Users\amiru\Desktop\Healthcare_MEDTHRU_Proposal.docx'

doc = Document(SRC)


# --------------------------------------------------------------------------
# block helpers
# --------------------------------------------------------------------------
def iter_blocks(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)


BLOCKS = list(iter_blocks(doc))
PARAS = [b for b in BLOCKS if isinstance(b, Paragraph)]
TABLES = [b for b in BLOCKS if isinstance(b, Table)]

used = set()


def find(prefix, blocks=None):
    """First unused paragraph whose text starts with `prefix`."""
    for i, p in enumerate(blocks or PARAS):
        if i in used:
            continue
        if p.text.strip().startswith(prefix):
            used.add(i)
            return p
    raise LookupError('paragraph not found: ' + prefix[:60])


def set_text(p, value, style=None):
    """Replace a paragraph's text, keeping run 0's character formatting.

    `value` is a string, or a list of (text, bold) pairs for bold lead-ins.
    """
    parts = [(value, None)] if isinstance(value, str) else list(value)
    runs = p.runs
    if not runs:
        r0 = p.add_run('')
        runs = p.runs
    r0 = runs[0]
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    r0.text = parts[0][0]
    if parts[0][1] is not None:
        r0.bold = parts[0][1]
    for text, bold in parts[1:]:
        r = p.add_run(text)
        r.font.name = r0.font.name
        r.font.size = r0.font.size
        r.bold = bool(bold)
    if style:
        p.style = doc.styles[style]
    return p


def drop(p):
    p._p.getparent().remove(p._p)


def add_after(anchor, value, style):
    """Insert a new paragraph directly after `anchor`, cloned for formatting."""
    new_p = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p)
    p = Paragraph(new_p, anchor._parent)
    for r in list(p.runs)[1:]:
        r._r.getparent().remove(r._r)
    p.style = doc.styles[style]
    set_text(p, value)
    return p


def chain_after(anchor, items, style):
    """Insert several paragraphs in order after `anchor`."""
    prev = anchor
    for it in items:
        prev = add_after(prev, it, style)
    return prev


def cell_text(cell, value):
    """Replace a table cell's text, keeping the first paragraph's formatting."""
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    set_text(cell.paragraphs[0], value)


def fill_table(tbl, rows, skip_header=True):
    """rows: list of row-lists of cell strings; None leaves a cell alone."""
    body = tbl.rows[1:] if skip_header else tbl.rows
    for row, values in zip(body, rows):
        for cell, value in zip(row.cells, values):
            if value is not None:
                cell_text(cell, value)


def drop_row(tbl, idx):
    tbl._tbl.remove(tbl.rows[idx]._tr)


def drop_spacer_after(tbl):
    """Remove the empty spacer paragraph that follows a table."""
    nxt = tbl._tbl.getnext()
    if nxt is not None and nxt.tag == qn('w:p') and not Paragraph(nxt, tbl._parent).text.strip():
        nxt.getparent().remove(nxt)


# --------------------------------------------------------------------------
# TEAM MEMBERS
# --------------------------------------------------------------------------
set_text(find('Team diversity'),
         'Team diversity: two women and two men. Two members meet the Women in '
         'Engineering (WIE) requirement.')

# --------------------------------------------------------------------------
# 1. EXECUTIVE SUMMARY   (must stay under 250 words)
# --------------------------------------------------------------------------
EXEC = [
    "When a patient arrives at an emergency department unable to speak, the "
    "information needed to treat them safely often already exists. The problem "
    "is that nobody can reach it in the first few minutes. Malaysian health "
    "records are split between public hospitals, private hospitals and general "
    "practices, and these systems do not share data. Phone based emergency "
    "profiles do not solve this, because they need a phone that is "
    "present, charged and unlocked.",

    "MEDTHRU proposes the Medical Identification Card (MedIC). It joins a "
    "patient mobile app, an NFC card and a clinical portal into one workflow. "
    "The card stores no medical data, only a random 256 bit token. The "
    "server keeps only the SHA-256 hash of that token. A valid tap shows a small "
    "emergency dataset, such as blood group, severe allergies, high risk "
    "medications and emergency contacts. Full history needs a clinician login, "
    "and every read and write is saved to an audit log. A lost card can be "
    "cancelled and replaced without changing the record.",

    "MedIC also uses AI as bounded decision support. It turns 30 days of vital "
    "sign readings into three short points. It also checks a small, clinician "
    "reviewed set of rules for dangerous drug combinations. The AI does not "
    "diagnose, prescribe or block any clinical decision. The build phase will "
    "finish the hardware loop, the AI subsystem and the hardened emergency "
    "route. The cost stays under RM 200 per station and under RM 3 per card.",
]
set_text(find('When an incapacitated patient arrives'), EXEC[0])
set_text(find('The card stores a pseudonymous'), EXEC[1])
set_text(find('MedIC also applies AI'), EXEC[2])

# --------------------------------------------------------------------------
# 2. BACKGROUND AND PROBLEM STATEMENT
# --------------------------------------------------------------------------
set_text(find('During an emergency, the first 60 minutes'),
         'The first 60 minutes after a serious injury are often called the '
         '"Golden Hour" (Lerner & Moscati, 2001). When a patient reaches an '
         'emergency department (ED) unconscious or in severe distress, staff may '
         'have no reliable identity or medical information. Time is then lost on '
         'identifying the patient, finding records, or asking worried family '
         'members. Three problems add up here:')

set_text(find('Delayed registration and identification'),
         [('Slow registration and identification. ', True),
          ('Manual typing and identity searches slow registration down, and that '
           'can delay assessment. In one Malaysian tertiary emergency department, '
           'crowding was linked to longer door to antibiotic times in sepsis '
           'patients (Chau et al., 2024). A delay at the front desk therefore '
           'becomes a delay in treatment.', False)])

set_text(find('Treatment without critical context'),
         [('Treatment without important context. ', True),
          ('Doctors often have to decide quickly, without confirmed information '
           'on severe allergies, high risk medications or chronic illness.', False)])

set_text(find('Inaccessible patient-held data'),
         [('Patient data that cannot be reached. ', True),
          ('Health data kept on a personal phone is useless if the patient is '
           'unconscious, the phone is locked, or there is no signal.', False)])

set_text(find('The problem is structural rather than incidental'),
         'This problem is structural, not accidental. Malaysian health records '
         'sit in three separate places. Government facilities use Ministry of '
         'Health systems, private hospitals use their own vendor platforms, and '
         'general practitioners use paper or standalone software. A patient may '
         'see a GP in Kajang, then a specialist in Cheras, and later arrive at '
         'Hospital Kuala Lumpur. That one patient can end up as three records '
         'that are not linked (Gong et al., 2024). No national summary travels '
         'with the patient. The people who suffer most are also the people that '
         'app only solutions serve worst. Older adults, rural communities and '
         'lower income patients often have complex histories. They are also the '
         'least likely to arrive with a charged phone, mobile data and a '
         'remembered password (Institute for Public Health, 2024).')

# 2.3 limitations table: smartphone apps and cloud portals merged, 6 rows to 5
drop_row(TABLES[1], 4)
fill_table(TABLES[1], [
    ['National EMR integration',
     'Needs agreement between institutions and years of procurement. It matters '
     'at national scale, but a small team cannot deploy it.'],
    ['Smartphone health apps and cloud patient portals',
     'Assume the patient is awake, has a charged phone, remembers the password '
     'and has a signal. All four can fail in an emergency. They also need a '
     'connection at the point of care.'],
    ['Paper cards and medical alert bracelets',
     'Works when the patient is unconscious, but holds only a few fixed fields, '
     'goes out of date, and cannot be cancelled if it is lost.'],
    ['National ID cards with medical fields (e.g. MyKad)',
     'The chip can hold blood type, allergies and chronic conditions. However, '
     'only Telehealth Flagship facilities can enter the data, the fields are '
     'fixed, and the holder cannot update or cancel them (Jabatan Pendaftaran '
     'Negara, n.d.).'],
    ['Phone based emergency profiles (Apple Medical ID, Android Emergency Information)',
     'Readable from a locked screen, but still needs a phone that is present and '
     'charged. The data is self reported, and staff must retype it into the '
     'hospital system.'],
])

set_text(find('The unmet need is therefore'),
         'What is missing is a record that travels with the patient, works '
         'without the patient\u2019s help, stays up to date and can be cancelled. '
         'In short, it needs the availability of a paper card and the freshness '
         'of a live database.')

# 2.4 stakeholder table
fill_table(TABLES[2], [
    ['Patients and cardholders',
     'A simple way to record daily health data, and confidence that key medical '
     'facts can still be found when they cannot speak',
     'If they cannot speak, they may be treated without staff knowing about '
     'severe allergies or existing conditions'],
    ['ED registration and triage staff',
     'Fast confirmation of who the patient is, with less repeated typing',
     'Slow systems, typing errors and record searches hold the team up during a rush'],
    ['Emergency doctors and specialists',
     'Direct access to a reliable summary of history, conditions, medications '
     'and recent trends',
     'Decisions rest on incomplete records, or on what a stressed family member '
     'can remember'],
])

# --------------------------------------------------------------------------
# 3. OBJECTIVES
# --------------------------------------------------------------------------
set_text(find('The project is assessed against seven'),
         'We will judge the project against seven measurable objectives. Each one '
         'states a target and how we will check it during the build phase.')

fill_table(TABLES[3], [
    ['O1', 'Remove identification delay at triage',
     'A card tap fills in the patient identity in under 5 seconds, measured over '
     '30 taps in a row on a cold started terminal.'],
    ['O2', 'Give emergency data without a login',
     'Layer 1 opens on card possession in 100% of valid taps. The API rejects '
     '100% of Layer 2 requests that carry no clinician credential.'],
    ['O3', 'Keep patient records separate',
     'No card ever returns another patient\u2019s rows. Checked by automated '
     'tests on every token based route.'],
    ['O4', 'Contain a lost card',
     'Cancelling and reissuing finish in one transaction, and the cancelled '
     'token then returns nothing.'],
    ['O5', 'Produce useful AI output',
     'A three point summary from 30 days of readings, and a visible warning in '
     '100% of the scripted blood thinner cases.'],
    ['O6', 'Keep the cost low',
     'Under RM 200 per station and under RM 3 per card, with staff training '
     'under 30 minutes.'],
    ['O7', 'Keep full accountability',
     'Every read and write, including patient self access, creates an audit row. '
     'Verified by test.'],
])

# --------------------------------------------------------------------------
# 4. PROPOSED SOLUTION
# --------------------------------------------------------------------------
set_text(find('The Medical Identification Card is an integrated'),
         'The Medical Identification Card is a three part platform. It links '
         'patients, triage staff and treating doctors through one workflow. '
         'Instead of separate apps and manual typing, it joins data entry, card '
         'based access and clinical review together.')

set_text(find('Patient health application.'),
         [('Patient health application. ', True),
          ('Patients record vital signs, manage chronic conditions and keep their '
           'medication list up to date. The app turns daily entries into simple '
           'trend charts, so a doctor can see patterns such as blood pressure over '
           'the past 30 days. Home monitoring of this kind has been linked to '
           'better blood pressure results (Harrison et al., 2025).', False)])

set_text(find('NFC access card.'),
         [('NFC access card. ', True),
          ('One tap reads a token and loads the patient\u2019s basic profile on the '
           'hospital screen in under five seconds. This cuts repeated typing, so '
           'assessment can start sooner.', False)])

set_text(find('Clinical portal.'),
         [('Clinical portal. ', True),
          ('Triage staff register the patient here, and doctors review the record. '
           'Access follows the two layer model in section 4.4.', False)])

set_text(find('Writing medical data directly to the card'),
         'Writing medical data straight onto the card creates three problems. The '
         'data goes out of date unless every card is rewritten by hand. Anyone who '
         'finds the card can read the record. A normal NTAG chip is also too small '
         'to hold a useful medical history.')

set_text(find('Instead, the platform writes a 256-bit'),
         'Instead, we write a random 256 bit token into the card\u2019s user data '
         'blocks. The token is a key, not a patient number. The database stores '
         'only its SHA-256 hash. The raw token is written once when the card is '
         'issued. After that it is only sent during a tap, hashed and compared, '
         'and never saved on the server. This gives four benefits at once:')

set_text(find('Patient presence provides emergency access'),
         'The card alone gives emergency access. An unconscious patient is '
         'identified by holding the card, with no password or phone.')
set_text(find('Information remains current'),
         'The data stays current. The card points to a live record instead of a '
         'fixed copy.')
set_text(find('Loss can be contained'),
         'A lost card can be contained. Cancelling stops the token working, and '
         'reissuing creates a new token in the same transaction.')
set_text(find('The card stores no identifying'),
         'The card holds nothing readable. It carries only a token. That token '
         'still counts as personal data in law, because it leads to an '
         'identifiable patient, and we treat it that way.')

set_text(find('Every NFC tag broadcasts a factory-assigned UID'),
         'Every NFC tag broadcasts a factory UID. Using it as a credential would '
         'be convenient, but it is not safe. The UID is sent to any nearby reader '
         'and can be copied with cheap hardware. Anyone who walks past a patient '
         'could then pretend to be that card later. We store our token in the '
         'tag\u2019s data blocks instead, and use the UID only for logging. This is '
         'the most important security decision in the design, and it is already '
         'implemented in reader.js and tag-storage.js.')

# 4.4 heading is styled Normal in the source
set_text(find('4.4 Dual-Layer Data Access'),
         '4.4 Two-Layer Data Access and Privacy Model', style='Heading 2')

set_text(find('A dual-layer architecture balances'),
         'The two layer model balances fast access to life saving data against '
         'protection of sensitive records. Layer 1 opens right after a valid tap, '
         'with no clinician login. It is limited to data that can change immediate '
         'treatment: verified blood group, severe allergies, active high risk '
         'medications, key chronic diagnoses and emergency contacts. The blood '
         'group is shown for context, not as the sole basis for a transfusion. '
         'Layer 2 holds detailed history, daily trends and diagnostic records. It '
         'needs a clinician login, and each access writes an audit entry with the '
         'clinician\u2019s identity, the time and the purpose. The API enforces '
         'this, not the interface. A client cannot obtain data that the server is '
         'not allowed to return.')

fill_table(TABLES[4], [
    ['First responder or triage staff', 'Layer 1 emergency dataset', 'Card possession'],
    ['Patient using own card', 'Layer 2, own record only',
     'Card possession, plus a card password if enabled'],
    ['Patient without the card', 'Layer 2, own record only',
     'Registered phone number and password'],
    ['Clinician', 'Layer 2, plus permission to write clinical data',
     'Clinician login (JWT)'],
])

set_text(find('The portal allows clinicians to review'),
         'Doctors can review and digitally verify what a patient has entered, so '
         'emergency staff can tell verified entries apart from self reported ones. '
         'The source attribution is already built. Every clinical row carries a '
         'source field, and the server sets it according to whether the request '
         'came with a valid clinician token. A patient entry is never shown as '
         'doctor verified. The build phase adds a stamp that records which doctor '
         'checked the entry and when.')

set_text(find('The platform operates alongside existing hospital'),
         'The platform runs beside existing hospital systems instead of replacing '
         'them. A facility keeps its current medical record system and only adds a '
         'USB reader at each triage station. The card then works as a patient held '
         'index across facilities that do not otherwise share data. The portal '
         'passes the basic registration fields through a REST endpoint, with an '
         'audited clipboard or CSV fallback for older systems. Clinical history is '
         'never included in that fallback. Cards are issued at reception after '
         'staff check the patient\u2019s MyKad and record consent. The facility '
         'acts as data controller, and our team acts as data processor under a '
         'written agreement.')

fill_table(TABLES[5], [
    ['1. Single clinic',
     'One reader at reception and one at the consultation desk, with cards issued '
     'at registration',
     'Tap to record works in a real workflow with real staff'],
    ['2. Clinic cluster', 'Several GP clinics sharing one backend',
     'Records follow the patient between facilities'],
    ['3. Emergency department', 'Readers at triage, Layer 1 only at the front desk',
     'The unconscious patient case, which is the highest value scenario'],
    ['4. Standards alignment', 'HL7 FHIR import and export over the existing schema',
     'A path into national infrastructure instead of another silo'],
])

# --------------------------------------------------------------------------
# 5. AI INTEGRATION STRATEGY
# --------------------------------------------------------------------------
set_text(find('During emergency care, clinicians have limited time'),
         'In an emergency, doctors have little time to read a long history. The '
         'platform runs text and time series analysis over the data collected by '
         'the patient app. This runs on the backend, not on the card, because the '
         'card only holds the token.')

set_text(find('Clinical executive summary.'),
         [('Clinical executive summary. ', True),
          ('The AI turns 30 days of vital sign data into three short points. It '
           'highlights patterns that are hard to spot in a raw table, such as blood '
           'pressure spikes or unstable blood sugar.', False)])

set_text(find('Triage risk and contraindication system.'),
         [('Triage risk and contraindication check. ', True),
          ('The system compares the patient\u2019s recorded allergies and current '
           'medications against the treatment being considered. For example, if a '
           'patient already takes a high dose blood thinner and the doctor plans to '
           'give more, the portal shows a clear warning. Warnings like these work '
           'best when they are specific and easy to act on (Holbrook et al., 2025). '
           'The aim is to help doctors notice serious risks, not to decide for '
           'them.', False)])

set_text(find('5.2 Clinical Safety'),
         '5.2 Clinical Safety and Human Oversight', style='Heading 2')

set_text(find('These functions provide decision support'),
         'These functions are decision support, not automatic decisions. Every '
         'summary and warning is advisory, carries the name of its source, and '
         'appears next to the raw data it came from. The system never hides data '
         'and never blocks a clinical action. It shows the risk, and it records '
         'that the warning was displayed. We chose this on purpose. A 2024 meta '
         'analysis found that clinicians ignore about 90% of drug interaction '
         'alerts in decision support systems (Felisberto et al., 2024). The '
         'platform therefore raises only warnings that are specific to the patient '
         'and worth acting on.')

anchor = set_text(find('AI is also integrated into prototype design'),
                  'We also use AI in the design, development and testing of the '
                  'prototype. The Nexus Log records this work as it happens. Our '
                  'main uses are:')
chain_after(anchor, [
    [('Architecture decisions. ', True),
     ('The token versus UID analysis, the two layer access model and the source '
      'attribution scheme came out of AI assisted design review. The team '
      'checked, challenged and sometimes rejected the suggestions.', False)],
    [('Protocol work. ', True),
     ('NTAG21x block read and write sequencing, and PC/SC command handling for '
      'the ACR122U reader.', False)],
    [('Security review. ', True),
     ('Adversarial analysis of the token lifecycle, the rate limits and the lost '
      'card threat model.', False)],
    [('Testing. ', True),
     ('Endpoint and access scoping tests, including the rule that card A must '
      'never reach patient B\u2019s rows.', False)],
], 'List Bullet')

# --------------------------------------------------------------------------
# 6. SYSTEM DESIGN
# --------------------------------------------------------------------------
set_text(find('The system follows a secure client'),
         'The system uses a secure client and server design. It has NFC hardware, '
         'cross platform clients, a clinical portal and a backend API. A patient '
         'records health data in the mobile app, and the backend links that record '
         'to the token on their card. At the hospital, the card is tapped on a USB '
         'reader at the triage terminal, and the token is sent to the backend over '
         'an encrypted connection. The backend checks the token and returns the '
         'Layer 1 profile. If a doctor needs the full history, logging in unlocks '
         'Layer 2 and starts the AI pipeline. All traffic is protected with TLS.')

set_text(find('Deployment model and its limitation'),
         [('Deployment model and its limits. ', True),
          ('The prototype keeps the backend inside the facility. The API and '
           'database sit on an ordinary PC or a hospital virtual machine, and the '
           'reader bridge reaches them over the local network. This keeps the '
           'record inside the facility that holds it. It does not make the system '
           'network free, because a tap still has to reach the local backend. '
           'Section 8.2 lists two planned fixes: an offline Layer 1 cache at each '
           'reader station, and a read only copy on a second host.', False)])

set_text(find('The cold-start path matters'),
         'Cold start matters for the demonstration and for real use. A tap works '
         'even when the terminal is idle. If the backend is not running, the '
         'reader bridge starts it. If the client is not open, the bridge opens it. '
         'The only thing the operator has to do is present the card.')

set_text(find('The card is a passive 13.56 MHz'),
         'The card is a passive 13.56 MHz tag. It draws power from the '
         'reader\u2019s field, has no battery, and needs no maintenance. The '
         'ACR122U connects to Windows as a standard PC/SC device and runs on USB '
         'power, so a triage station needs no extra adapter. The doctor\u2019s '
         'terminal needs no reader of its own, because it reaches the same record '
         'over the local network. Adding a review station therefore costs nothing '
         'beyond the PC already on the desk.')

set_text(find('Figure 3.'),
         'Figure 3. Physical layout of the card, reader and triage terminal. The '
         'doctor\u2019s terminal connects over the local network.')

# 6.4 technology stack
fill_table(TABLES[6], [
    ['Hardware', 'ACR122U USB NFC reader (13.56 MHz, PC/SC) and NTAG215 cards',
     'Identifies the patient and reads the token from the tag'],
    ['Reader bridge', 'Node.js, nfc-pcsc, custom NTAG block I/O',
     'Detects taps, reads the token, tells the clients, and cold starts the stack'],
    ['Clients', 'Flutter and Dart. Patient app on Android, clinical portal on '
     'desktop and later on the web',
     'Daily logging, trend charts, appointments, triage check in and doctor review'],
    ['Backend API', 'Node.js, Express 5, REST and JSON',
     'Token lookup, access layer enforcement, CRUD and audit logging'],
    ['Security', 'bcrypt, separate patient and clinician JWTs, SHA-256 token '
     'hashing, per route rate limits, TLS',
     'Login, credential protection and resistance to guessing attacks'],
    ['Database', 'SQLite (16 tables) and a filesystem store for documents',
     'Clinical record, card lifecycle and append only audit trail'],
    ['AI subsystem', 'Time series anomaly detection, text summarisation, and a '
     'bounded rule based contraindication engine',
     'Turns long term data into useful insight at the point of care'],
])

set_text(find('Sixteen SQLite tables underpin'),
         'Sixteen SQLite tables support the platform. Figure 4 shows the security '
         'critical links. Each patient has one active card, and reissuing cancels '
         'the previous card in the same transaction. Every read and write is '
         'audited, including when patients view their own record. Routes that take '
         'a token find the patient first and then filter the child rows. One card '
         'therefore cannot reach another patient\u2019s record, even with a valid '
         'row number.')

# 6.5 threat table: 9 data rows down to 7
threats = TABLES[7]
drop_row(threats, 6)   # "Unaudited access", folded into the row above
drop_row(threats, 7)   # "Data at rest unencrypted", folded into the JWT row
fill_table(threats, [
    ['Copying the broadcast hardware UID',
     'The UID is not used as a credential. The token sits in the user data blocks '
     'instead.', 'Implemented'],
    ['Leak of the token database',
     'Only SHA-256 hashes are stored. The raw token is never saved on the server.',
     'Implemented'],
    ['Card lost or stolen',
     'Cancel and reissue. An optional card password protects the full record on a '
     'plain tap.', 'Implemented'],
    ['Guessing tokens, or a client granting itself access',
     'A 256 bit token space and 30 lookups per minute. Both layers are enforced on '
     'the server, so the interface is not a security boundary.', 'Implemented'],
    ['Reaching another patient\u2019s data, or access that leaves no record',
     'Every token based query is filtered by the resolved patient_id, and every '
     'read and write writes an audit_log row.', 'Implemented'],
    ['A copied static token from an NTAG215 card',
     'A prototype limit. Token entropy, cancellation, rate limits, audit '
     'monitoring and facility controlled readers reduce the risk. Production needs '
     'a tag that supports challenge and response.', 'Partly mitigated'],
    ['Hardcoded development JWT secret, and data at rest that is not encrypted',
     'The secret must come from an environment variable, and the database must be '
     'encrypted, before any real patient data is used.', 'Open'],
])

set_text(find('The final two items are stated deliberately'),
         'We state the open items on purpose. They are known blockers, and we must '
         'close them before the system handles real patient data.')

anchor = set_text(find('The core record platform is substantially implemented'),
                  'Most of the record platform is already built:')
last = chain_after(anchor, [
    [('Backend. ', True),
     ('About 1,750 lines of Node and Express across 16 SQLite tables, with full '
      'audit logging and CRUD for records, appointments, messaging and '
      'documents.', False)],
    [('Security. ', True),
     ('bcrypt password hashing, with separate clinician and patient JWTs.', False)],
    [('Reader bridge. ', True),
     ('A working PC/SC bridge with cold start and live tap broadcast, plus card '
      'provisioning for blank NTAG tags.', False)],
    [('Client. ', True),
     ('A 46 file Flutter app with trend charts, PDF export and notifications.', False)],
    [('Testing. ', True),
     ('A harness that runs the full tap path without hardware, and an endpoint '
      'test suite covering cross patient scoping.', False)],
], 'List Bullet')
add_after(last,
          'The build phase delivers the three remaining pieces: the physical '
          'hardware loop, the AI subsystem and the hardened Layer 1 route.',
          'Normal')

# --------------------------------------------------------------------------
# 7. BUDGET
# --------------------------------------------------------------------------
fill_table(TABLES[8], [
    ['ACR122U USB NFC reader', '1', '126', '126',
     'The PC/SC reader needed to close the card to reader loop (O1). One unit '
     'equips the demonstration triage station, and it lasts about five years.'],
    ['NTAG215 blank cards', '3', '2.33', '7',
     'Patient cards for the demonstration, including one spare to show the cancel '
     'and reissue flow (O4).'],
    ['NTAG213 sticker labels', '10', '0.60', '6',
     'Cheap tags for repeated read and write testing, so we do not use up real '
     'patient cards.'],
    ['Total', '', '', '139', 'Confirms the target of under RM 200 per station in O6.'],
])

# --------------------------------------------------------------------------
# 8. IMPROVEMENT PLAN
# --------------------------------------------------------------------------
fill_table(TABLES[9], [
    ['16 to 22 Sep', 'Close the hardware loop from end to end',
     'A cold boot tap shows the Layer 1 profile in under 5 seconds, unattended'],
    ['18 to 25 Sep', 'AI summary over 30 days of readings',
     'A three point summary from real logged data, with unusual values flagged'],
    ['23 to 28 Sep', 'Contraindication engine, Layer 1 hardening, lost card flow',
     'A blood thinner test case raises a visible warning, and a cancelled card is '
     'refused'],
    ['26 to 30 Sep', 'Booth build, demo script, cards provisioned',
     'The demonstration runs three times in a row without help'],
    ['30 Sep to 7 Oct', 'Documentation, source code, video and Nexus Log',
     'All four items complete and submitted at least 24 hours early'],
])
# keeps section 9 on the same page as the rest of section 8
drop_spacer_after(TABLES[9])

set_text(find('Immediate.'),
         [('Immediate. ', True),
          ('Fix the JWT secret, add encryption at rest, and run a security review. '
           'Make the audit trail tamper evident with hash chaining, so changes made '
           'by a database administrator can be detected. Replace the desktop client '
           'with a web portal, add HL7 FHIR export, and build an offline cache for '
           'clinics with weak connectivity.', False)])
set_text(find('Medium-term.'),
         [('Medium term. ', True),
          ('Use Android NFC so that a phone becomes the reader, which drops the '
           'hardware cost per station to zero. Add hospital single sign on, so '
           'clinician identity comes from the facility directory. Extend the AI to '
           'triage scoring and early warning of deterioration.', False)])
set_text(find('Long-term.'),
         [('Long term. ', True),
          ('Let facilities federate, so hospitals look up a card against a shared '
           'directory while each one keeps custody of its own records. Run a '
           'Ministry of Health pilot in an emergency department.', False)])

set_text(find('Success will be evaluated through four'),
         'We will measure success with four operational measures, not download '
         'counts. First, the time from arrival to confirmed identity at triage. '
         'Second, the time until critical history is available. Third, how much '
         'doctors trust the AI summary when they check it against the record. '
         'Fourth, the share of issued cards still active after 12 months, which '
         'shows whether patients keep carrying the card.')

# --------------------------------------------------------------------------
# 9. CONCLUSION
# --------------------------------------------------------------------------
set_text(find('MedIC closes the emergency information gap'),
         'MedIC closes the emergency information gap with three parts. These are '
         'a patient health app, an NFC card that carries a token instead of '
         'medical data, and a portal with two layer access. The build phase '
         'delivers the hardware loop, the AI subsystem and the hardened Layer 1 '
         'route.')

# --------------------------------------------------------------------------
# References: one em dash in the NHMS entry
# --------------------------------------------------------------------------
ref = find('Institute for Public Health. (2024)')
set_text(ref, ref.text.replace('\u2014', ':').replace('demand::', 'demand:'))

# --------------------------------------------------------------------------
# Drop the hand-inserted page breaks. They were tuned against the old, longer
# wording and now strand most of a page each (section 5 started a fresh page
# with only 3 inches used on the page before it). The keep_with_next and
# keep_together rules already stop headings and figures splitting, and
# References keeps its own pageBreakBefore.
# --------------------------------------------------------------------------
removed = 0
for p in list(doc.paragraphs):
    brs = [b for b in p._p.iter(qn('w:br')) if b.get(qn('w:type')) == 'page']
    if brs and not p.text.strip():
        drop(p)
        removed += 1

# --------------------------------------------------------------------------
# Figures 1 to 3 are rendered far above print resolution (585 to 854 dpi at
# their current size). Scaling them to 85% recovers about an inch of vertical
# space and still leaves every one of them above 680 dpi.
# --------------------------------------------------------------------------
NSW = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
NSA = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
SCALE = 0.85
figure_n = 0
for p in doc.paragraphs:
    for r in p.runs:
        for inl in r._r.iter(NSW + 'inline'):
            figure_n += 1
            if figure_n == 1 or figure_n > 4:   # 1 is the cover, 5 is Figure 4
                continue
            for e in inl.iter():
                if e.tag in (NSW + 'extent', NSA + 'ext') and e.get('cx'):
                    e.set('cx', str(int(int(e.get('cx')) * SCALE)))
                    e.set('cy', str(int(int(e.get('cy')) * SCALE)))

# --------------------------------------------------------------------------
# safety net: no em dash anywhere in the body
# --------------------------------------------------------------------------
for block in iter_blocks(doc):
    if isinstance(block, Paragraph):
        targets = [block]
    else:
        targets = [p for row in block.rows for c in row.cells for p in c.paragraphs]
    for p in targets:
        for r in p.runs:
            if '\u2014' in r.text:
                r.text = re.sub(r'\s*\u2014\s*', ', ', r.text)

doc.save(OUT)

words = sum(len(p.split()) for p in EXEC)
print('Manual page breaks removed:', removed)
print('Executive summary word count:', words)
print('Saved:', OUT)
